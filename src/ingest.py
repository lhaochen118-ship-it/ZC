#!/usr/bin/env python3
"""Generic ingestion of one business data package (CSV / XLSX / SQLite) into clean pandas tables + a compact data card.

Nothing here is keyed on a business id: header rows, number formats ("10.000đ", "105k", "1,5tr", "3 ly"), mixed date
formats and multi-file / multi-sheet splits of one list are detected from the data itself.

load_business(dir) -> Biz(tables: {name: DataFrame (typed)}, raw: {name: DataFrame (str)}, meta: {name: info}, readme)
Biz.card()         -> text description of every table for the LLM prompt
Usage (cache all businesses of a split): ingest.py <split_dir> <cache_dir> [n_workers]"""
import os, re, sys, csv, glob, json, pickle, sqlite3, unicodedata, datetime as dt
from collections import OrderedDict
import numpy as np
import pandas as pd

# ---------------------------------------------------------------- value parsing
_CUR = re.compile(r"(vnđ|vnd|đồng|đ|₫|d)$")
_NUM = re.compile(r"^[-+]?\(?\d[\d.,\s]*\)?")


def parse_num(x, dec=None):
    """Vietnamese-style number -> float (nan if not a number). '10.000đ'->10000, '8,000.0000'->8000, '105k'->105000,
    '1,5tr'->1500000, '3 ly'->3, '(1.000)'->-1000, '18.6'->18.6. `dec` ('.' or ',') is the decimal separator inferred for
    the whole column (see dec_mode); without it a single separator followed by exactly 3 digits means thousands."""
    if x is None: return np.nan
    if isinstance(x, (bool, np.bool_)): return np.nan
    if isinstance(x, (int, float, np.integer, np.floating)): return float(x)
    s = unicodedata.normalize("NFC", str(x)).strip().lower().replace(" ", " ")
    if not s or s in ("nan", "none", "null", "-", "—"): return np.nan
    s = re.sub(r"^(vnđ|vnd|đ|₫|usd|\$)\s*", "", s)                       # currency written before the number
    m = re.fullmatch(r"(\d+)\s*(tr|triệu|trieu|k|nghìn|nghin|ngàn|ngan|tỷ|ty)\s*(\d{1,3})\s*(đ|vnđ|vnd|₫)?", s)
    if m:                                                              # shorthand money: 1tr440 = 1,440,000 ; 5k5 = 5,500 ; 2tr9 = 2,900,000
        mult = 1e6 if m.group(2).startswith("tr") else 1e3 if m.group(2)[0] in "kn" else 1e9
        return float(m.group(1)) * mult + float(m.group(3)) * mult / 10 ** len(m.group(3))
    m = _NUM.match(s)
    if not m: return np.nan
    tok, rest = m.group(0).strip(), s[m.end():].strip()
    if re.fullmatch(r"\d{1,3}(\s\d{3})+", tok): tok = tok.replace(" ", "")   # '531 500' (space thousands)
    elif re.search(r"\d\s+\d", tok):                                    # '1 xe 4m3' is fine below, '1 6 2025' is not a number
        tok, rest = tok.split()[0], s[len(tok.split()[0]):].strip()
    if re.match(r"^[/:\-.]\s*\d", rest) or re.match(r"^(thg|tháng|th)\s*\d", rest): return np.nan   # dates / times
    if re.match(r"^(tr|triệu|trieu|k|nghìn|nghin|ngàn|ngan|tỷ|ty)", rest) and tok.count(",") + tok.count(".") == 1:
        dec = "," if "," in tok else "."                                 # '1,134tr' = 1.134 million: before a multiplier the separator is decimal
    neg = tok.startswith("-") or (tok.startswith("(") and tok.endswith(")"))
    tok = tok.strip("+-() ").replace(" ", "")
    amb = re.fullmatch(r"\d{1,3}[.,]\d{3}", tok) is not None           # '1.500' / '1,500': thousands or decimal?
    if "," in tok and "." in tok:
        d_ = "," if tok.rfind(",") > tok.rfind(".") else "."
        tok = tok.replace("." if d_ == "," else ",", "").replace(d_, ".")
    elif "," in tok:
        if tok.count(",") > 1 or (amb and dec != ","): tok = tok.replace(",", "")
        else: tok = tok.replace(",", ".")
    elif "." in tok:
        if tok.count(".") > 1 or (amb and dec != "."): tok = tok.replace(".", "")
    try: v = float(tok)
    except ValueError: return np.nan
    rest = _CUR.sub("", rest).strip()
    # multipliers only as whole words: "1 kg" is one kilogram (not 1,000), "10 m" ten metres (not 10 million)
    if re.match(r"(tr|triệu|trieu)(?![a-zà-ỹ])", rest): v *= 1e6
    elif re.match(r"(k|nghìn|nghin|ngàn|ngan)(?![a-zà-ỹ])", rest): v *= 1e3
    elif re.match(r"(tỷ|ty)(?![a-zà-ỹ])", rest): v *= 1e9
    v = round(v, 6)                                                    # '65,079tr' -> 65079000, not 65078999.99999999
    return -v if neg else v


_DATE_PATS = [(re.compile(r"^(\d{1,2})[/\-.](\d{1,2})[/\-.](\d{4})(?:[ T](\d{1,2}):(\d{2})(?::(\d{2}))?)?$"), "dmy"),
              (re.compile(r"^(\d{4})[/\-.](\d{1,2})[/\-.](\d{1,2})(?:[ T](\d{1,2}):(\d{2})(?::(\d{2})(?:\.\d+)?)?)?$"), "ymd"),
              (re.compile(r"^(\d{1,2})[/\-.](\d{1,2})[/\-.](\d{2})$"), "dmy2"),
              (re.compile(r"^(?:ngày\s*)?(\d{1,2})\s*(?:thg|tháng|th)\s*(\d{1,2})[,\s]*(?:năm\s*)?(\d{4})(?:[ T,]+(\d{1,2}):(\d{2})(?::(\d{2}))?)?$", re.I), "dmy")]
_MON = {m: i + 1 for i, m in enumerate(["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"])}


def parse_date(x, year=None):
    """-> pd.Timestamp or NaT. dd/mm/yyyy is assumed for slash dates (Vietnamese convention). `year` completes 'dd/mm'."""
    if x is None: return pd.NaT
    if isinstance(x, (pd.Timestamp, dt.datetime)): return pd.Timestamp(x)
    if isinstance(x, dt.date): return pd.Timestamp(x)
    if isinstance(x, (int, float, np.integer, np.floating)):
        return pd.Timestamp("1899-12-30") + pd.Timedelta(days=float(x)) if 20000 < x < 80000 else pd.NaT   # Excel serial
    s = str(x).strip()
    for pat, kind in _DATE_PATS:
        m = pat.match(s)
        if not m: continue
        g = [int(v) if v else 0 for v in m.groups()]
        try:
            if kind == "dmy": return pd.Timestamp(g[2], g[1], g[0], g[3], g[4], g[5])
            if kind == "ymd": return pd.Timestamp(g[0], g[1], g[2], g[3], g[4], g[5])
            if kind == "dmy2": return pd.Timestamp(2000 + g[2], g[1], g[0])
        except ValueError: return pd.NaT
    m = re.fullmatch(r"(\d{1,2})[/\-.](\d{1,2})", s)                     # dd/mm without year
    if m and year:
        try: return pd.Timestamp(int(year), int(m.group(2)), int(m.group(1)))
        except ValueError: return pd.NaT
    m = re.match(r"^(\d{1,2})[\s\-/]([a-z]{3})[a-z]*[\s\-/,]+(\d{4})", s.lower())          # 01-Jan-2025
    if m and m.group(2) in _MON:
        try: return pd.Timestamp(int(m.group(3)), _MON[m.group(2)], int(m.group(1)))
        except ValueError: return pd.NaT
    return pd.NaT


def norm_text(s):
    """Scorer-style normalisation (NFC, lower, punctuation -> space, collapse)."""
    s = unicodedata.normalize("NFC", str(s)).lower()
    return re.sub(r"\s+", " ", "".join(c if c.isalnum() or c.isspace() else " " for c in s)).strip()


def strip_accents(s):
    s = unicodedata.normalize("NFD", str(s)).replace("đ", "d").replace("Đ", "D")
    return unicodedata.normalize("NFC", "".join(c for c in s if unicodedata.category(c) != "Mn"))


# ---------------------------------------------------------------- reading files
def _rows_xlsx(path):
    try:
        from python_calamine import CalamineWorkbook                     # fast reader if available
        wb = CalamineWorkbook.from_path(path)
        for name in wb.sheet_names: yield name, wb.get_sheet_by_name(name).to_python(skip_empty_area=False)
        return
    except ImportError: pass
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    for ws in wb.worksheets: yield ws.title, [list(r) for r in ws.iter_rows(values_only=True)]


def _rows_csv(path):
    raw = open(path, "rb").read()
    for enc in ("utf-8-sig", "utf-16", "cp1258", "latin-1"):
        try: txt = raw.decode(enc); break
        except UnicodeDecodeError: continue
    try: dialect = csv.Sniffer().sniff(txt[:20000], delimiters=",;\t|")
    except csv.Error: dialect = csv.excel
    yield "", [r for r in csv.reader(txt.splitlines(), dialect)]


def _blank(v): return v is None or (isinstance(v, float) and np.isnan(v)) or (isinstance(v, str) and not v.strip())


def _find_header(rows):
    """Index of the header row: first row (within 30) with >= 60% of the widest row's cells filled, mostly text."""
    head = rows[:30]
    width = max((sum(not _blank(v) for v in r) for r in head), default=0)
    for i, r in enumerate(head):
        cells = [v for v in r if not _blank(v)]
        if width and len(cells) >= max(2, 0.6 * width) and sum(isinstance(v, str) and np.isnan(parse_num(v)) for v in cells) >= 0.6 * len(cells):
            return i
    return 0


def _to_frame(rows):
    rows = [list(r) for r in rows]
    while rows and all(_blank(v) for v in rows[-1]): rows.pop()
    if not rows: return None, []
    h = _find_header(rows)
    meta = [" ".join(str(v) for v in r if not _blank(v)) for r in rows[:h]]
    meta = [m for m in meta if m]
    header = list(rows[h]); body = rows[h + 1:]
    names = [str(v).strip() for v in header if not _blank(v)]
    if h > 0 and len(names) != len(set(names)):                        # duplicated sub-headers: use the group row above
        grp = list(rows[h - 1]) + [None] * (len(header) - len(rows[h - 1])); cur = None; filled = []
        for j in range(len(header)):
            if not _blank(grp[j]): cur = str(grp[j]).strip()
            filled.append(cur)
        if sum(not _blank(g) for g in grp) >= 2:
            header = [f"{filled[j]} - {header[j]}" if filled[j] and not _blank(header[j]) else header[j] for j in range(len(header))]
            meta = meta[:-1] if meta and meta[-1] == " ".join(str(v) for v in grp if not _blank(v)) else meta
    width = max(len(header), max((len(r) for r in body), default=0))
    cols, seen = [], {}
    for j in range(width):
        c = header[j] if j < len(header) and not _blank(header[j]) else f"col{j + 1}"
        c = re.sub(r"\s+", " ", str(c)).strip()
        if c in seen: seen[c] += 1; c = f"{c}.{seen[c]}"
        else: seen[c] = 0
        cols.append(c)
    hdr = {str(v).strip() for v in header if not _blank(v)}
    def is_header(r):                                                  # header repeated inside the body (page breaks, concatenated exports)
        cells = [str(v).strip() for v in r if not _blank(v)]
        return len(cells) >= 2 and sum(c in hdr for c in cells) >= 0.6 * len(cells)
    def is_note(r):                                                    # 'Tổng cộng', 'Ghi chú: ...' rows: sparse and start with a summary keyword
        cells = [str(v).strip() for v in r if not _blank(v)]
        return bool(cells) and len(cells) <= max(1, 0.4 * width) and any(_NOTE.match(c) for c in cells)
    kept = [r + [None] * (width - len(r)) for r in body if not all(_blank(v) for v in r) and not is_header(r)]
    notes = [" ".join(str(v) for v in r if not _blank(v))[:80] for r in kept if is_note(r)]
    body = [r for r in kept if not is_note(r)]
    meta += [f"[dropped {len(notes)} summary/note rows, e.g. {notes[:2]}]"] if notes else []
    df = pd.DataFrame(body, columns=cols, dtype=object)
    df = df[[c for c in cols if not df[c].map(_blank).all()]]                                # drop all-empty columns
    return df, meta


_NOTE = re.compile(r"^(tổng|tong|cộng|cong|total|sum|ghi chú|ghi chu|note|lưu ý|luu y|người lập|nguoi lap|ký tên)\b", re.I)


def _base_name(fname, sheet):
    stem = os.path.splitext(os.path.basename(fname))[0]
    stem = re.sub(r"([_\-\s]*(T?\d{1,2}[_\-]?)?\d{4}|[_\-\s]\d{1,3})$", "", stem) or stem   # BanHang_T1_2025 -> BanHang
    return stem if not sheet else f"{stem}"


# ---------------------------------------------------------------- typing
_CODEY = re.compile(r"(^|[\s_(])(mã|ma|code|id|sđt|sdt|đt|dt|phone|điện thoại|dien thoai|tel|số hđ|so hd|barcode|mst)([\s_).]|$)", re.I)


def _num_tail(v):
    """Text left after the leading number, without currency words ('' for plain numbers)."""
    if not isinstance(v, str): return ""
    s = unicodedata.normalize("NFC", v).strip().lower()
    m = _NUM.match(s)
    if not m: return ""
    tok = m.group(0).strip()
    if re.search(r"\d\s+\d", tok) and not re.fullmatch(r"\d{1,3}(\s\d{3})+", tok): tok = tok.split()[0]
    rest = s[s.index(tok) + len(tok):].strip() if tok in s else ""
    rest = re.sub(r"^(tr|triệu|trieu|k|nghìn|nghin|ngàn|ngan|tỷ|ty)\d*\b", "", rest).strip()
    return _CUR.sub("", rest).strip(" .%")


def dec_mode(strs):
    """Decimal separator of a column from its own values: '.' if some values are clearly '.'-decimals (1.5, 2.25,
    0.1234) and none are clearly '.'-thousands (1.234.567, 1.234,5); symmetric for ','. None if no evidence."""
    money = strs.str.contains(r"\d\s*(?:đ|vnđ|vnd|₫|k|tr|triệu|nghìn|ngàn)\b|^(?:vnd|vnđ)", regex=True).mean()
    if money >= 0.05: return None                                      # money columns: per-value rule (3 digits after a separator = thousands)
    st = strs[strs.str.fullmatch(r"-?\d[\d.,]*") & ~strs.str.fullmatch(r"-?\d+[.,]0+")]   # plain numbers; '37000.0' is a float artefact
    dd = st.str.fullmatch(r"-?\d+\.(\d{1,2}|\d{4,})").sum(); cd = st.str.fullmatch(r"-?\d+,(\d{1,2}|\d{4,})").sum()
    dt_ = st.str.fullmatch(r"-?\d{1,3}(\.\d{3}){2,}").sum() + st.str.contains(r"\d\.\d{3},\d", regex=True).sum()
    ct = st.str.fullmatch(r"-?\d{1,3}(,\d{3}){2,}").sum() + st.str.contains(r"\d,\d{3}\.\d", regex=True).sum()
    if dd and not dt_ and not cd: return "."
    if cd and not ct and not dd: return ","
    return None


def auto_type(df, year_hint=None):
    """Type every column from its values only (no column-name lists except generic code/address words):
    - datetime if >= 90% of values parse as dates (datetime cells, date strings, Excel serials mixed with date strings,
      'dd/mm' completed with the year found in the sheet / file / title of the row);
    - number if >= 95% parse as numbers; text after the number is kept in `<col>__unit` when it looks like a unit
      (few distinct short tails), otherwise the column stays text (addresses, '12 Nguyễn Trãi');
    - long digit strings (codes, IMEI, phone numbers, leading zeros) stay text."""
    out = df.copy(); types = {}
    yh = year_hint if year_hint is not None else pd.Series([None] * len(df), index=df.index)
    for c in df.columns:
        if c.startswith("_"): continue
        s = df[c]; isb = s.map(_blank); nn = s[~isb]
        if len(nn) == 0: types[c] = "empty"; continue
        idx = nn.sample(min(len(nn), 500), random_state=0).index if len(nn) > 500 else nn.index
        smp = nn.loc[idx]
        is_dt = smp.map(lambda v: isinstance(v, (pd.Timestamp, dt.datetime, dt.date)))
        is_str = smp.map(lambda v: isinstance(v, str))
        def pdate(v, y):                                               # inside a date column, '46183' / 46183 are Excel serials
            if isinstance(v, str) and re.fullmatch(r"\d{5}(\.\d+)?", v.strip()): v = float(v)
            return parse_date(v, y)
        dates = pd.Series([pdate(v, yh.get(i)) for i, v in smp.items()], index=idx)
        date_evidence = (is_dt | smp.map(lambda v: isinstance(v, str) and pd.notna(parse_date(v, 2000)))).mean() >= 0.2
        if dates.notna().mean() >= 0.9 and date_evidence:
            out[c] = pd.Series([pd.NaT if b else pdate(v, yh.get(i)) for (i, v), b in zip(s.items(), isb)], index=s.index); types[c] = "datetime"; continue
        strs = smp.astype(str).str.strip()
        codey = _CODEY.search(c) or strs.str.match(r"^0\d{3,}$").any() or strs.str.fullmatch(r"\d{12,}").mean() > 0.5
        if not codey and not _ADDR.search(c):
            dm = dec_mode(nn.astype(str).str.strip().str.lower())
            nums = smp.map(lambda v: parse_num(v, dm))
            if nums.notna().mean() >= 0.95:
                tails = smp.map(_num_tail); tl = tails[tails != ""]
                unit_like = len(tl) == 0 or (tl.nunique() <= max(30, 0.02 * len(smp)) and tl.str.split().str.len().max() <= 3)
                if unit_like:
                    out[c] = s.map(lambda v: np.nan if _blank(v) else parse_num(v, dm)).astype(float); types[c] = "number"
                    if len(tl):
                        out[c + "__unit"] = s.map(lambda v: _num_tail(v) or None); types[c + "__unit"] = "text"
                    continue
        out[c] = s.map(lambda v: None if _blank(v) else (unicodedata.normalize("NFC", str(v)).strip() if not isinstance(v, (int, float)) else v))
        types[c] = "text"
    cols = []                                                          # keep each __unit column next to its number column
    for c in df.columns:
        cols.append(c)
        if c + "__unit" in out.columns: cols.append(c + "__unit")
    return out[cols], types


_ADDR = re.compile(r"(địa chỉ|dia chi|địa điểm|address|addr|đường|duong)", re.I)


# ---------------------------------------------------------------- business
class Biz:
    def __init__(self, path):
        self.path = path; self.tables = OrderedDict(); self.raw = OrderedDict(); self.meta = OrderedDict()
        rd = os.path.join(path, "README.md")
        self.readme = open(rd, encoding="utf-8").read() if os.path.isfile(rd) else ""

    def card(self, max_examples=4):
        return "\n".join(table_card(name, df, self.meta[name]["types"], self.meta[name]["sources"], self.meta[name].get("title_rows"),
                                     self.raw[name], max_examples) for name, df in self.tables.items())


def col_type(s):
    if pd.api.types.is_datetime64_any_dtype(s): return "datetime"
    if pd.api.types.is_bool_dtype(s): return "bool"
    if pd.api.types.is_numeric_dtype(s): return "number"
    return "text"


def table_card(name, df, types=None, sources=None, title_rows=None, raw=None, max_examples=4):
    """Compact description of one DataFrame for the LLM prompt."""
    L = [f"### TABLE `{name}`  rows={len(df)}" + (f"  from: {', '.join(sources[:6])}{' …' if len(sources) > 6 else ''}" if sources else "")]
    if title_rows: L.append(f"title rows above header (dropped): {' | '.join(title_rows[:3])[:200]}")
    for c in df.columns:
        s = df[c]; t = (types or {}).get(c) or col_type(s); nn = s.notna().mean() if len(s) else 0
        try:
            if t == "number": d = f"min={s.min():g} max={s.max():g}"
            elif t == "datetime": d = f"{s.min()} .. {s.max()}"
            elif t == "bool": d = f"true={int(s.fillna(False).astype(bool).sum())}"
            else:
                vc = s.dropna().astype(str); nu = vc.nunique()
                ex = list(dict.fromkeys(vc.head(2000)))[:max_examples] if nu > 12 else list(vc.value_counts().index)
                d = f"n_unique={nu} " + ("values=" if nu <= 12 else "e.g.=") + json.dumps([e[:40] for e in ex], ensure_ascii=False)
        except Exception as e: d = f"({type(e).__name__})"
        L.append(f"- `{c}` [{t}{'' if nn > 0.999 else f', {nn:.0%} filled'}] {d}")
    if raw is not None:
        L.append("raw first row: " + json.dumps(raw.head(1).to_dict("records"), ensure_ascii=False, default=str)[:400])
    return "\n".join(L)


def load_business(path):
    biz = Biz(path); groups = OrderedDict()
    files = sorted(f for f in glob.glob(os.path.join(path, "*")) if not f.endswith(".md"))
    for f in files:
        ext = os.path.splitext(f)[1].lower()
        if ext in (".sqlite", ".db", ".sqlite3"):
            con = sqlite3.connect(f"file:{f}?mode=ro", uri=True)
            for (t,) in con.execute("select name from sqlite_master where type='table'"):
                df = pd.read_sql(f'select * from "{t}"', con)
                groups.setdefault(("sqlite", t, tuple(df.columns)), []).append((f"{os.path.basename(f)}:{t}", df.astype(object), [], t))
            con.close(); continue
        reader = _rows_xlsx if ext in (".xlsx", ".xlsm", ".xls") else _rows_csv
        for sheet, rows in reader(f):
            df, title = _to_frame(rows)
            if df is None or df.empty: continue
            key = (_base_name(f, sheet), tuple(df.columns))
            src = os.path.basename(f) + (f"[{sheet}]" if sheet else "")
            groups.setdefault(key, []).append((src, df, title, sheet))
    names = {}
    for key, parts in groups.items():
        name = key[1] if key[0] == "sqlite" else key[0]
        if name in names: names[name] += 1; name = f"{name}_{names[name]}"
        else: names[name] = 0
        frames, hints = [], []
        for src, df, title, sheet in parts:
            df = df.copy(); df.insert(0, "_source", src)
            if sheet and len(parts) > 1: df.insert(1, "_sheet", sheet)
            yrs = [re.search(r"(?<!\d)(19|20)\d\d(?!\d)", t) for t in [str(sheet or ""), src] + list(title)]
            hints += [next((int(m.group(0)) for m in yrs if m), None)] * len(df)
            frames.append(df)
        raw = pd.concat(frames, ignore_index=True)
        typed, types = auto_type(raw, pd.Series(hints, index=raw.index))
        biz.raw[name] = raw; biz.tables[name] = typed
        biz.meta[name] = {"sources": [p[0] for p in parts], "types": types, "title_rows": parts[0][2]}
    return biz


def cache_split(split_dir, cache_dir, workers=4):
    from concurrent.futures import ProcessPoolExecutor
    os.makedirs(cache_dir, exist_ok=True)
    bids = sorted(os.listdir(os.path.join(split_dir, "data")))
    def todo(b): return not os.path.isfile(os.path.join(cache_dir, b + ".pkl"))
    with ProcessPoolExecutor(workers) as ex:
        for b, msg in zip([b for b in bids if todo(b)], ex.map(_cache_one, [(split_dir, cache_dir, b) for b in bids if todo(b)])):
            print(b, msg, flush=True)


def _cache_one(a):
    split_dir, cache_dir, b = a
    import time; t = time.time()
    biz = load_business(os.path.join(split_dir, "data", b))
    pickle.dump(biz, open(os.path.join(cache_dir, b + ".pkl"), "wb"), protocol=4)
    open(os.path.join(cache_dir, b + ".card.md"), "w", encoding="utf-8").write(biz.card())
    return f"{len(biz.tables)} tables {sum(len(d) for d in biz.tables.values())} rows {time.time() - t:.0f}s"


if __name__ == "__main__":
    import ingest                                                      # pickle Biz as ingest.Biz, not __main__.Biz
    ingest.cache_split(sys.argv[1], sys.argv[2], int(sys.argv[3]) if len(sys.argv) > 3 else 4)
